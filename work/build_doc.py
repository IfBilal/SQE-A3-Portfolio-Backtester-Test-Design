import json, os, re, glob, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
import plan, plan_bva, content

ROOT='/home/bilaltahir/Desktop/SQE-A3'
S=json.load(open(ROOT+'/work/schedule.json'))
FORMNAME={f[0]:f[1] for f in plan.FORMS}
FORMURL={f[0]:f[2] for f in plan.FORMS}
ORDER=[f[0] for f in plan.FORMS]

# ---------- manifest ----------
MAN={}
extra=ROOT+'/work/manifest_extra.txt'
for fn in [ROOT+'/manifest.txt', extra]:
    if os.path.exists(fn):
        for l in open(fn, encoding='utf-8'):
            l=l.rstrip('\n')
            if not l.strip(): continue
            p=l.split('|',2)
            if len(p)<3: continue
            MAN[p[0]]=(p[1], p[2])
# image mapping
IMG={}
for f in glob.glob(ROOT+'/shots/*.jpg'):
    b=os.path.basename(f)
    if b.startswith('s2_'): IMG[b[:-4]]=f
    else:
        n=b[:-4].split('-')[-1]; IMG['n'+n]=f; IMG[n]=f
def imgs(tc):
    if tc not in MAN: return []
    out=[]
    for k in MAN[tc][0].split(','):
        k=k.strip()
        if k in IMG: out.append(IMG[k])
    return out

# ---------- verdict ----------
FAILS={'TC-EP-S3','TC-EP-S4','TC-EP-S6'}
PASSES={'TC-EP-S2'}
def verdict(s):
    tc=s['id']
    if tc not in MAN: return None
    res=MAN[tc][1]
    mode=s['mode']
    if tc in FAILS: return 'Fail'
    if tc in PASSES: return 'Pass'
    if res.startswith('PLAN-LIMIT') and mode=='gate' and s['form'] in ('MC','EF'): return 'Pass'
    if s['form'] in ('LD','RB'):
        return 'Blocked'
    if mode in ('run','run_any'):
        return 'Pass' if (res.startswith('RAN') or res.startswith('ACCEPTED')) else 'Fail'
    if mode in ('native','toast','run_bad'):
        return 'Pass' if res.startswith('REJECTED') else 'Fail'
    if mode=='gate': return 'Pass' if res.startswith('PLAN-LIMIT') else 'Fail'
    if mode=='pending':
        ex=s['exp'].lower()
        if any(w in ex for w in ('rejected','limited','required')):
            return 'Pass' if res.startswith('REJECTED') else 'Fail'
        return 'Pass' if (res.startswith('RAN') or res.startswith('ACCEPTED')) else 'Fail'
    return None
def actual(s):
    tc=s['id']
    if tc not in MAN: return 'Not executed'
    r=MAN[tc][1].strip()
    parts=[p.strip() for p in r.split(' | ') if p.strip()]
    parts=[p for p in parts if not re.match(r'^toast: Backtest( optimization)? complete$',p,re.I) and not re.match(r'^toast: Withdrawal policy analysis complete$',p,re.I)]
    out=' | '.join(parts)
    out=out.replace('REJECTED (native browser validation + toast)','Rejected (browser validation + toast)').replace('REJECTED (native browser validation)','Rejected (browser validation message)').replace('REJECTED (toast)','Rejected (toast)')
    return out
VERD={s['id']:verdict(s) for s in S}

# ---------- docx helpers ----------
doc=Document()
sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for m in ('left_margin','right_margin'): setattr(sec,m,Cm(1.5))
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.5)
st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
for h,sz in (('Heading 1',18),('Heading 2',14),('Heading 3',11.5)):
    doc.styles[h].font.name='Calibri'; doc.styles[h].font.size=Pt(sz); doc.styles[h].font.color.rgb=RGBColor(0x1F,0x3A,0x5F)

def hdr_footer():
    h=sec.header.paragraphs[0]; h.text='24i-3166  |  Muhammad Bilal Tahir  |  SE-B  |  SE 3002 - Assignment 3: Black-Box Testing (EP & BVA) on Testfolio'
    h.runs[0].font.size=Pt(8.5); h.runs[0].font.bold=True
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=f.add_run('Roll No: 24i-3166   |   Page '); r.font.size=Pt(8.5)
    for t,txt in (('begin',None),('instr',' PAGE '),('end',None)):
        run=f.add_run(); run.font.size=Pt(8.5)
        if t=='instr':
            e=OxmlElement('w:instrText'); e.set(qn('xml:space'),'preserve'); e.text=txt
        else:
            e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),t)
        run._r.append(e)
hdr_footer()

def shade(cell, color):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),color); tcPr.append(sh)
def set_cell(cell, text, bold=False, size=8, color=None, align=None):
    cell.text=''
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
    r=p.add_run(str(text)); r.font.size=Pt(size); r.bold=bold
    if color: r.font.color.rgb=RGBColor.from_string(color)
    if align: p.alignment=align
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); e=OxmlElement('w:tblHeader'); e.set(qn('w:val'),'true'); trPr.append(e)
def no_split(row):
    trPr=row._tr.get_or_add_trPr(); e=OxmlElement('w:cantSplit'); e.set(qn('w:val'),'true'); trPr.append(e)
def table(headers, rows, widths, size=8, vcol=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; set_cell(c,h,True,size,'FFFFFF'); shade(c,'1F3A5F'); c.width=Cm(widths[i])
    repeat_header(t.rows[0])
    for r in rows:
        row=t.add_row(); no_split(row)
        for i,v in enumerate(r):
            c=row.cells[i]; c.width=Cm(widths[i])
            col=None; b=False
            if vcol is not None and i==vcol:
                if v=='Pass': col='1B7F3B'; b=True
                elif v=='Fail': col='C0392B'; b=True
                elif v=='Blocked': col='B9770E'; b=True
            set_cell(c,v,b,size,col)
    for i,w in enumerate(widths): t.columns[i].width=Cm(w)
    return t
def para(text, bold=False, size=10, italic=False, after=4):
    p=doc.add_paragraph(); r=p.add_run(text); r.bold=bold; r.italic=italic; r.font.size=Pt(size); p.paragraph_format.space_after=Pt(after); return p
def bullets(items, size=10):
    for it in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2)
        r=p.add_run(it); r.font.size=Pt(size)
def kv_table(rows):
    t=doc.add_table(rows=0, cols=2); t.style='Table Grid'
    for k,v in rows:
        r=t.add_row(); set_cell(r.cells[0],k,True,10); shade(r.cells[0],'E8EEF5'); set_cell(r.cells[1],v,False,10)
        r.cells[0].width=Cm(7); r.cells[1].width=Cm(19)
    t.columns[0].width=Cm(7); t.columns[1].width=Cm(19)
    return t

CFG=json.load(open(ROOT+'/work/student.json')) if os.path.exists(ROOT+'/work/student.json') else {}
ACCOUNT=CFG.get('account_email','(Testfolio account email - to be entered)')

# ---------- cover ----------
doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('SE 3002 - Software Quality Engineering'); r.bold=True; r.font.size=Pt(20)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('FAST School of Computing - Department of Software Engineering - FAST NUCES Islamabad').font.size=Pt(11)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('ASSIGNMENT 3 - Student Submission'); r.bold=True; r.font.size=Pt(18)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Black-Box Testing: Equivalence Partitioning & Boundary Value Analysis on Testfolio (testfol.io)').font.size=Pt(13)
doc.add_paragraph()
kv_table([('Student Name','Muhammad Bilal Tahir'),('Roll Number','24i-3166'),('Section / Batch','SE-B'),
          ('Instructors','Dr. Nigar Azhar Butt  |  Dr. Atif Aftab Jilani  |  Miss Anum Kaleem'),
          ('Test Account Used',ACCOUNT),('Browser & Version','Google Chrome 142.0.0.0 on Ubuntu Linux')])
doc.add_paragraph()
n_ex=sum(1 for s in S if s['id'] in MAN)
para('Summary of this submission: 11 forms, %d input fields, %d equivalence classes, %d boundaries, %d EP + Strong EP test cases and %d BVA test cases (%d executed with screenshots, %d marked Not Executed); %d defect observations.'
     %(len(plan.FIELDS), sum(len(r) for _,_,r in plan.FIELDS), len(plan_bva.BC), sum(1 for s in S if s['kind'] in ('EP','EPS')), sum(1 for s in S if s['kind']=='BVA'), n_ex, len(S)-n_ex, len(content.DEFECTS)), italic=True)
doc.add_page_break()

# ---------- D1 ----------
doc.add_heading('D1 - Test Plan [20 marks]',1)
doc.add_heading('TP-1  Test Plan Metadata',2)
kv_table([('Test Plan ID',content.PLAN['id']),('Version',content.PLAN['version']),('Author (Roll No.)',content.PLAN['author']),('Status',content.PLAN['status'])])
doc.add_heading('TP-2  Introduction',2); para(content.PLAN['intro'])
doc.add_heading('TP-3  Scope',2); para('In-scope features (all 11 forms and their input fields):'); bullets(content.PLAN['scope'])
doc.add_heading('TP-4  Test Objectives',2); bullets(content.PLAN['objectives'])
doc.add_heading('TP-5  Test Approach',2)
for a in content.PLAN['approach']: para(a)
para('Designated critical form: Portfolio Backtester (testfol.io/). Justification: it has the largest number of inputs, the most validation rules, cross-field rules (date order, allocation total) and the highest risk (money values and date logic). It receives Strong EP and 3-point BVA.', bold=True)
doc.add_heading('TP-6  Test Environment',2)
e=content.PLAN['env']
kv_table([('Browser & Version',e['browser']),('Operating System',e['os']),('Screen Resolution',e['res']),('Test Account (email)',ACCOUNT),('Network / VPN',e['net']),('Plan',e['plan'])])
doc.add_heading('TP-7  Entry & Exit Criteria',2); para('Entry criteria:',True); bullets(content.PLAN['entry']); para('Exit criteria:',True); bullets(content.PLAN['exit'])
doc.add_heading('TP-8  Test Schedule',2)
table(['Activity','Notes / Dates','Est. Hours','Done?'], [[a,b,c,'Yes'] for a,b,c in content.PLAN['schedule']], [9,10,3,3], 9)
doc.add_heading('TP-9  Risks & Mitigations',2)
table(['Risk','Likelihood / Impact','Mitigation'], content.PLAN['risks'], [11,4,11], 9)
doc.add_page_break()
doc.add_heading('TP-10  Probe Log - Appendix A',2)
para('Every probe lists all inputs entered (context fields included). Probes were run on my own account before any test case was designed.', italic=True, size=9)
table(['Probe ID','Form','Test Inputs (field: value, ...)','System Response','Inference / Note'],
      [['PR-%03d'%(i+1),a,b,c,d] for i,(a,b,c,d) in enumerate(content.PROBES)], [1.8,3.2,8.5,6.5,6.3], 8)
extra_probes=ROOT+'/work/probes_extra.json'
doc.add_page_break()

# ---------- D2 ----------
doc.add_heading('D2 - Equivalence Partitioning [25 marks]',1)
doc.add_heading('Layer 1 - EP Partition Tables (one per input field)',2)
FL=sorted(plan.FIELDS,key=lambda x:ORDER.index(x[0]))
ec=0; ECMAP={}
for form,field,rows in FL:
    doc.add_heading('%s - %s'%(FORMNAME[form],field),3)
    trs=[]
    for (desc,rep,typ,exp,sets,mode) in rows:
        ec+=1; ECMAP[(form,field,desc)]='EC-%03d'%ec
        trs.append(['EC-%03d'%ec,desc,rep,typ])
    table(['EC ID','Equivalence Class Description','Representative Value','Type (VEC / IEC)'],trs,[2.2,12,8,3.5],8.5)
doc.add_page_break()

doc.add_heading('Layer 2 - EP Test Case Table (Weak EP with robustness: one TC per EC)',2)
ep=[s for s in S if s['kind']=='EP']
rows=[]
for s in ep:
    rows.append([s['id'],s['ec'],FORMNAME[s['form']],s['field'],s['rep'],s['exp'],actual(s),VERD[s['id']] or 'Not Executed'])
table(['TC-EP ID','EC IDs Covered','Feature / Form','Input Field(s)','Test Value(s)','Expected Result','Actual Result','P/F'],rows,[1.9,1.7,3.0,3.2,3.8,4.2,6.7,1.6],7,7)
doc.add_page_break()
doc.add_heading('Strong EP - Critical form: Portfolio Backtester (several classes covered in one test case)',2)
sp=[s for s in S if s['kind']=='EPS']
# strong EP cover which ECs: describe
STRONGEC={'TC-EP-S1':'All valid classes: start/end dates, Starting Value, Rolling Window, growth, cashflow, drag, ticker and allocation','TC-EP-S2':'IEC of Starting Value + IEC of Rolling Window + IEC of cashflow growth','TC-EP-S3':'IEC unknown ticker + IEC allocation sum below 100','TC-EP-S4':'IEC start after end + IEC empty Starting Value + IEC unknown ticker + IEC allocation sum above 100','TC-EP-S5':'VEC dates and numerics + IEC negative allocation','TC-EP-S6':'IEC empty ticker + IEC drag above 100'}
table(['TC-EP ID','EC IDs Covered','Feature / Form','Input Field(s)','Test Value(s)','Expected Result','Actual Result','P/F'],
      [[s['id'],STRONGEC[s['id']],'Portfolio Backtester','multiple fields',s['rep'],s['exp'],actual(s),VERD[s['id']] or 'Not Executed'] for s in sp],[1.9,5.0,2.6,2.2,4.2,3.6,5.4,1.2],7,7)
doc.add_page_break()

# ---------- D3 ----------
doc.add_heading('D3 - Boundary Value Analysis [25 marks]',1)
para('Judgment on field selection: BVA was applied to fields with numeric, length or ordering boundaries discovered by probing. Fields excluded: Ticker, Email and Password format, Rebalance select (symbolic input, no boundary); Cashflow amount (no limit found); date fields on forms other than the ordering rule on the critical form.', italic=True, size=9)
doc.add_heading('Layer 1 - BVA Boundary Tables',2)
BCL=sorted(plan_bva.BC,key=lambda x:ORDER.index(x[0]))
# find EC id for BC: use ecref text -> look up by field & matching class heuristic
def ecid_for(form,field,ecref):
    keys=[k for k in ECMAP if k[0]==form and (k[1]==field or field.startswith(k[1]) or k[1].startswith(field.split(' / ')[0]))]
    pick=[]
    for k in keys:
        if k[2] in ecref or ecref.split(' (')[0].split(' vs ')[0].lower() in k[2].lower():
            pick.append(ECMAP[k])
    return ', '.join(pick[:2]) if pick else (ECMAP[keys[0]] if keys else '-')
bc_rows={}
cur=None
for bi,(form,field,ecref,desc,bval,cases) in enumerate(BCL,1):
    bc_rows.setdefault((form,field),[]).append(['BC-%03d'%bi,FORMNAME[form],field,ecid_for(form,field,ecref),desc,bval])
for (form,field),rows in bc_rows.items():
    doc.add_heading('%s - %s'%(FORMNAME[form],field),3)
    table(['BC ID','Feature / Form','Input Field','EC ID','Boundary Description','Boundary Value'],rows,[2,4.5,4.5,4,8.5,3.2],8.5)
doc.add_page_break()
doc.add_heading('Layer 2 - BVA Test Case Table (2-point on all forms; 3-point on the critical form)',2)
bva=[s for s in S if s['kind']=='BVA']
rows=[[s['id'],s['bc'],FORMNAME[s['form']],s['field'],s['rep'],s['pos'],s['exp'],actual(s),VERD[s['id']] or 'Not Executed'] for s in bva]
table(['TC-BVA ID','BC IDs','Feature / Form','Input Field','Test Value','Position','Expected Result','Actual Result','P/F'],rows,[1.9,1.5,3.0,3.0,2.8,1.6,3.6,7.4,1.5],7,8)
doc.add_page_break()

# ---------- D4 ----------
doc.add_heading('D4 - Traceability Matrix [10 marks]',1)
para('Pass = behaviour matches the expected result; Fail = behaviour deviates (defect observation); Blocked = the Free plan redirects to the pricing page before the validation can be observed. Every test case from D2 and D3 is listed.', italic=True, size=9)
rows=[]
for s in S:
    tech='Strong EP' if s['kind']=='EPS' else ('Weak EP + robustness' if s['kind']=='EP' else ('3-point BVA' if s['form']=='BT' else '2-point BVA'))
    cls=(s['ec']+': '+s['desc']) if s['kind']=='EP' else ((s['bc']+' ('+s['pos']+'): '+s['desc']) if s['kind']=='BVA' else s['desc'])
    st_='Executed' if s['id'] in MAN else 'Not Executed'
    rows.append([s['id'],FORMNAME[s['form']],s['field'],tech,cls,st_,VERD[s['id']] or '-'])
table(['TC ID','Feature / Form','Input Field','Technique','Class / Boundary Description','Status','Result'],rows,[2,3.6,3.8,3.2,10.5,1.8,1.6],7,6)
doc.add_page_break()

# ---------- D5 ----------
doc.add_heading('D5 - Execution Log [20 marks]',1)
para('Screenshots (roll-number note visible) for every executed test case. The verdict, actual result and defect observations are in D2/D3; the defect list below groups them.', italic=True, size=9)
doc.add_heading('Defect observations',2)
table(['Defect','Form','Observation','Test cases'],content.DEFECTS,[1.6,4,14.5,6.5],8)
doc.add_page_break()
CACHE=ROOT+'/work/img_cache'; os.makedirs(CACHE,exist_ok=True)
def small(path):
    out=os.path.join(CACHE,os.path.basename(path))
    if not os.path.exists(out):
        im=Image.open(path).convert('RGB'); w=900
        if im.width>w: im=im.resize((w,int(im.height*w/im.width)),Image.LANCZOS)
        im.save(out,'JPEG',quality=55,optimize=True)
    return out
doc.add_heading('Screenshots by test case',2)
curform=None
for s in S:
    if s['id'] not in MAN: continue
    im=imgs(s['id'])
    if not im: continue
    if s['form']!=curform:
        curform=s['form']; doc.add_heading('%s (%s)'%(FORMNAME[curform],FORMURL[curform]),3)
    t=doc.add_table(rows=2, cols=max(1,len(im))); t.style='Table Grid'; t.autofit=False
    for i in range(len(im)): t.rows[0].cells[i].merge(t.rows[0].cells[i]) 
    cap=t.rows[0].cells[0]
    if len(im)>1: cap=cap.merge(t.rows[0].cells[len(im)-1])
    set_cell(cap,'%s  |  %s  |  input: %s  |  result: %s'%(s['id'],s['field'],s['rep'],(VERD[s['id']] or '')),True,8)
    cap.paragraphs[0].paragraph_format.keep_with_next=True
    shade(cap,'E8EEF5')
    for i,pth in enumerate(im):
        c=t.rows[1].cells[i]; c.width=Cm(13.2 if len(im)>1 else 24.8)
        c.paragraphs[0].paragraph_format.space_after=Pt(0)
        c.paragraphs[0].add_run().add_picture(small(pth), width=Cm(12.9 if len(im)>1 else 24.0))
    no_split(t.rows[0]); no_split(t.rows[1])
    sp_=doc.add_paragraph(); sp_.paragraph_format.space_after=Pt(2)
OUT=ROOT+'/A3_24i-3166_MuhammadBilalTahir'
doc.save(OUT+'.docx'); print('saved docx', os.path.getsize(OUT+'.docx')//1024, 'KB')
